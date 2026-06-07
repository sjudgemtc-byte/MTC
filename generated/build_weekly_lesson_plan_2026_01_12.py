from pathlib import Path

from docx import Document


BASE = Path("/Users/simonjudge/Documents/MTC")
TEMPLATE = BASE / "01_Lesson_Plans/lesson_plans/templates/ET F419 - Weekly Lesson Plan Master.docx"
OUTPUT = BASE / "01_Lesson_Plans/lesson_plans/examples/ET_F419_Lesson_Plan_2026_01_12_to_2026_01_16.docx"


def set_cell_text(cell, text):
    cell.text = text


def main():
    doc = Document(str(TEMPLATE))
    t0, t1, t2, t3 = doc.tables

    set_cell_text(t0.cell(2, 1), "Redfern")
    set_cell_text(t0.cell(2, 3), "RED-01 Pathway")
    set_cell_text(t0.cell(3, 1), "12/01/2026")
    set_cell_text(t0.cell(3, 3), "16/01/2026")
    set_cell_text(t0.cell(4, 1), "Simon Judge")
    set_cell_text(t0.cell(5, 1), "Only as needed during delivery.")
    set_cell_text(
        t0.cell(5, 3),
        "Use staged modelling, repeated instructions, login support, file-management support, reading support, and extension tasks where required.",
    )
    set_cell_text(
        t0.cell(6, 1),
        "☐ Mainstream SEE\n☒ Mainstream SEE (Canvas)\n☒ Workforce Preparation\n☐ Co-delivery (e.g. Business)\n☐ Industry Pathway (e.g. Pathway to Individual Care)\n☐ Community Access (e.g. Citizenship)\n☐ Distance\n☐ Other:",
    )
    set_cell_text(
        t0.cell(7, 1),
        "☐ 22637VIC Course in English as an Additional Language (Course in EAL)\n☒ 22688VIC Course in Initial General Education for Adults (CGEA Initial)\n☒ 22689VIC Certificate I General Education for Adults (Introductory) (CGEA I Intro)\n☐ FSK20119 Certificate II in Skills for Work and Vocational Pathways (FSK II)\n☐ Other:",
    )

    days = [
        {
            "name": "Monday",
            "date": "12 / 01 / 2026",
            "mode": "Synchronous  ☐\nAsynchronous  ☐\nNo Class Scheduled  ☒",
            "topic": "No class scheduled / preparation day",
            "units": "N/A - no scheduled class.",
            "objectives": "No scheduled class. Preparation only.",
            "staging": "Prepare digital literacy materials for Tuesday to Friday lessons.\nCheck Canvas course content, class files, and device-readiness requirements.\nOrganise differentiated support for learners needing login, navigation, and reading assistance.",
            "resources": "Previous lesson plan\nCanvas course materials\nPrepared classroom files\nDigital literacy resources",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Tuesday",
            "date": "13 / 01 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Digital literacy: creating a simple spreadsheet and accessing Canvas",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment",
            "objectives": "By the end of the day, learners will be able to open a spreadsheet program, create a simple table from a blank file, enter data into rows and columns, and access class content in Canvas with support.\nLearners will practise file naming, saving, and opening digital learning platforms.",
            "staging": "Session 1: 9:00-10:30 - teacher models opening a new spreadsheet, naming columns, and entering simple data using the Tokyo Uni Materials file as a guide.\nSession 2: 10:45-12:00 - learners create their own spreadsheet from scratch, enter provided data, and practise basic formula use with teacher support.\nSession 3: 12:30-13:45 - learners log in to the RED-01 Canvas course, locate assigned pages, and complete one guided online task linked to digital literacy or workplace learning.\nDifferentiation / support: one-to-one login support, shared-screen modelling, simplified data-entry version for beginners, and extension through extra rows or additional formulas.\nAssessment / evidence: created spreadsheet file, observed data entry and formula use, successful Canvas access, and teacher notes on digital independence.",
            "resources": "Tokyo Uni Materials.xlsx\nMicrosoft Excel or Google Sheets\nComputers or laptops with internet access\nRED-01 Canvas course\nTeacher model spreadsheet",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Wednesday",
            "date": "14 / 01 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Digital literacy and online reading: Canvas activities and habit stacking",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22350 Engage with short simple texts for learning purposes",
            "objectives": "By the end of the day, learners will be able to access assigned Canvas activities, read short on-screen content about habit stacking, and respond to guided digital tasks using supported reading strategies.\nLearners will practise navigating online learning content and selecting or typing responses with support.",
            "staging": "Session 1: 9:00-10:30 - teacher supports learners to log in to Canvas, locate assigned content, and review on-screen instructions.\nSession 2: 10:45-12:00 - teacher introduces the habit stacking summary and models how to read, scroll, and locate key information in a digital document or Canvas page.\nSession 3: 12:30-13:45 - learners complete follow-up Canvas activities, quizzes, or short typed responses linked to habit stacking and daily study routines.\nDifferentiation / support: read-aloud support, paired navigation help, short instructions, highlighted vocabulary, and extension through extra written detail or additional online activities.\nAssessment / evidence: accessed Canvas content, engaged with the habit stacking reading, completed digital responses or quiz items, and teacher observation of navigation and reading on screen.",
            "resources": "RED-01 Canvas course\nhabit_stacking_summary.docx\nComputers or laptops with internet access\nProjector or shared screen\nTeacher support notes",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Thursday",
            "date": "15 / 01 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Google Docs and Word processing: formatting a Wi-Fi factsheet and checking spelling",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22350 Engage with short simple texts for learning purposes",
            "objectives": "By the end of the day, learners will be able to open a digital document, apply simple formatting such as bold text and headings, use spell check tools with support, and read a short digital factsheet for key information.\nLearners will build confidence using basic word-processing tools for classroom and workplace tasks.",
            "staging": "Session 1: 9:00-10:30 - teacher models opening the Wi-Fi factsheet, identifying headings, and applying bold formatting and spell check tools.\nSession 2: 10:45-12:00 - learners complete the formatting task step by step, using teacher modelling to bold headings, correct spelling, and improve readability.\nSession 3: 12:30-13:45 - learners read selected sections of the factsheet, discuss key information, and complete a short follow-up digital response or comprehension task.\nDifferentiation / support: repeated demonstration, simplified checklist, shared reading, screenshot-style cues, and extension through extra formatting or extra corrections.\nAssessment / evidence: formatted document, observed use of spell check and toolbar features, short follow-up response, and teacher notes on digital task completion.",
            "resources": "factsheet-wifi.docx\nMicrosoft Word or Google Docs\nComputers or laptops\nProjector or shared screen\nTeacher formatting checklist",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Friday",
            "date": "16 / 01 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Digital literacy consolidation: Canvas tasks, online quizzes, and independent navigation",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment",
            "objectives": "By the end of the day, learners will be able to navigate the Canvas course more independently, complete assigned online activities or quizzes, and follow digital instructions to submit responses correctly.\nLearners will consolidate key digital skills from the week, including logging in, locating files, and completing online tasks.",
            "staging": "Session 1: 9:00-10:30 - teacher reviews the week's digital literacy skills, including spreadsheets, reading on screen, and word-processing tools.\nSession 2: 10:45-12:00 - learners work through assigned Canvas pages, quizzes, or surveys, applying navigation, reading, and response skills with teacher support as required.\nSession 3: 12:30-13:45 - learners finish online tasks, check completion, and review how to locate instructions and saved work; teacher provides catch-up support for learners needing additional help.\nDifferentiation / support: one-to-one support with navigation and instructions, repeated login support, paired help, and simplified completion targets for learners needing more time.\nAssessment / evidence: completed Canvas activity or quiz, observed navigation and task completion, and teacher notes on learner independence and support provided.",
            "resources": "RED-01 Canvas course\nComputers or laptops with internet access\nCanvas modules, quizzes, and surveys\nTeacher review checklist",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
    ]

    for col, day in enumerate(days):
        set_cell_text(t1.cell(0, col), f"{day['name']}\n{day['date']}")
        set_cell_text(t1.cell(1, col), day["mode"])
        set_cell_text(t1.cell(2, col), "Synchronous")
        set_cell_text(t1.cell(3, col), "Topic:")
        set_cell_text(t1.cell(4, col), day["topic"])
        set_cell_text(t1.cell(5, col), "Units of Competency:")
        set_cell_text(t1.cell(6, col), day["units"])
        set_cell_text(t1.cell(7, col), "Lesson objectives:")
        set_cell_text(t1.cell(8, col), day["objectives"])
        set_cell_text(t1.cell(9, col), "Staging/Activities:")
        set_cell_text(t1.cell(10, col), day["staging"])

        set_cell_text(t2.cell(0, col), "Resources:")
        set_cell_text(t2.cell(1, col), day["resources"])
        set_cell_text(t2.cell(2, col), "Asynchronous")
        set_cell_text(t2.cell(3, col), "Topic:")
        set_cell_text(t2.cell(4, col), day["async_topic"])
        set_cell_text(t2.cell(5, col), "Units of Competency:")
        set_cell_text(t2.cell(6, col), day["async_units"])
        set_cell_text(t2.cell(7, col), "Lesson objectives:")
        set_cell_text(t2.cell(8, col), day["async_objectives"])
        set_cell_text(t2.cell(9, col), "Staging/Activities:")
        set_cell_text(t2.cell(10, col), day["async_staging"])
        set_cell_text(t2.cell(11, col), "Resources:")
        set_cell_text(t2.cell(12, col), day["async_resources"])

    set_cell_text(
        t3.cell(1, 0),
        "Use short instructions, modelling, and repeated checks for learners needing support with reading, writing, and digital navigation.\n\n"
        "Provide repeated support with device login, Canvas access, opening files, saving work, and locating documents in Drive or shared folders.\n\n"
        "Use read-aloud support, highlighted keywords, sentence frames, and paired navigation for learners needing LLN support.\n\n"
        "Monitor confidence and engagement during online tasks and provide calm troubleshooting support where learners become stuck.\n\n"
        "Offer extension tasks for faster learners through additional formulas, extra formatting work, or more independent Canvas completion.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
