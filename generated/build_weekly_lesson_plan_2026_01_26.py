from pathlib import Path

from docx import Document


BASE = Path("/Users/simonjudge/Documents/MTC")
TEMPLATE = BASE / "01_Lesson_Plans/lesson_plans/templates/ET F419 - Weekly Lesson Plan Master.docx"
OUTPUT = BASE / "01_Lesson_Plans/lesson_plans/examples/ET_F419_Lesson_Plan_2026_01_26_to_2026_01_30.docx"


def set_cell_text(cell, text):
    cell.text = text


def main():
    doc = Document(str(TEMPLATE))
    t0, t1, t2, t3 = doc.tables

    set_cell_text(t0.cell(2, 1), "Redfern")
    set_cell_text(t0.cell(2, 3), "RED-01 Pathway")
    set_cell_text(t0.cell(3, 1), "26/01/2026")
    set_cell_text(t0.cell(3, 3), "30/01/2026")
    set_cell_text(t0.cell(4, 1), "Simon Judge")
    set_cell_text(t0.cell(5, 1), "Only as needed during delivery.")
    set_cell_text(
        t0.cell(5, 3),
        "Use staged modelling, repeated instructions, login support, file-management support, spreadsheet support, and extension tasks where required.",
    )
    set_cell_text(
        t0.cell(6, 1),
        "☐ Mainstream SEE\n☒ Mainstream SEE (Canvas)\n☒ Workforce Preparation\n☐ Co-delivery (e.g. Business)\n☐ Industry Pathway (e.g. Pathway to Individual Care)\n☐ Community Access (e.g. Citizenship)\n☐ Distance\n☐ Other:",
    )
    set_cell_text(
        t0.cell(7, 1),
        "☐ 22637VIC Course in English as an Additional Language (Course in EAL)\n☒ 22688VIC Course in Initial General Education for Adults (CGEA Initial)\n☒ 22689VIC Certificate I General Education for Adults (Introductory) (CGEA I Intro)\n☐ FSK20119 Certificate II in Skills for Work and Vocational Pathways (FSK II)\n☐ Other:",
    )

    days = [
        {
            "name": "Monday",
            "date": "26 / 01 / 2026",
            "mode": "Synchronous  ☐\nAsynchronous  ☐\nNo Class Scheduled  ☒",
            "topic": "No class scheduled / public holiday",
            "units": "N/A - no scheduled class.",
            "objectives": "No scheduled class.",
            "staging": "No scheduled class. Preparation and follow-up only if required.",
            "resources": "N/A",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Tuesday",
            "date": "27 / 01 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Digital setup and contact list skills",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22349 Create short simple texts for learning purposes\nVU22353 Use basic oral communication skills",
            "objectives": "By the end of the day, learners will be able to log in to classroom devices and Google tools with support, explain why digital contact lists are useful, and enter simple contact information into a structured table or document.\nLearners will practise typing, checking details, and using short spoken exchanges about names, phone numbers, and email addresses.",
            "staging": "Session 1: 9:00-10:30 - onboard learners, check computer access, batteries, Google login, Google Drive, and Gmail. Early finishers use typing.com or Be Connected.\nSession 2: 10:45-12:00 - introduce digital contact lists and discuss names, phone numbers, email addresses, and why accurate information matters in everyday digital life.\nSession 3: 12:30-13:45 - guided practice entering sample contact details into a simple digital table or document, followed by short oral review and class recap.\nDifferentiation / support: one-to-one login support, modelled text entry, vocabulary prompts, paired support, and extension through extra fields or categories.\nAssessment / evidence: successful login and setup, completed contact list task, observed typing and checking of details, and teacher notes on oral participation.",
            "resources": "Google account, Google Drive, and Gmail\nAddress book/contact list examples\nSpreadsheet or table template\ntyping.com\nBe Connected website\nComputers or laptops",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Wednesday",
            "date": "28 / 01 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Spreadsheet training: address book data entry and saving files",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment",
            "objectives": "By the end of the day, learners will be able to open the address book training spreadsheet, identify rows, columns, and cells, enter contact information accurately, and save and locate their work in Google Drive.\nLearners will build confidence using spreadsheet layout and basic data-entry conventions.",
            "staging": "Session 1: 9:00-10:30 - setup routine with login, Drive, Gmail, and independent typing.com or Be Connected practice for early finishers; teacher models the Address-Book-Training spreadsheet and key spreadsheet terms.\nSession 2: 10:45-12:00 - learners complete guided spreadsheet practice, entering names, phone numbers, email addresses, and categories into the correct cells.\nSession 3: 12:30-13:45 - whole-class recap, short oral questions about spreadsheet features, and review of saving and locating files in Google Drive.\nDifferentiation / support: shared-screen modelling, repeated explanation of rows and columns, one-to-one checking, and extension through extra entries or corrections.\nAssessment / evidence: completed spreadsheet entries, observed understanding of rows/columns/cells, saved file in Drive, and teacher notes on accuracy and independence.",
            "resources": "Address-Book-Training_2026_01_28.xlsx\nGoogle Sheets or Excel\nGoogle Drive\nProjector / shared screen\nComputers or laptops",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Thursday",
            "date": "29 / 01 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Digital setup and contact list editing skills",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22349 Create short simple texts for learning purposes\nVU22353 Use basic oral communication skills",
            "objectives": "By the end of the day, learners will be able to reopen their digital contact work, edit and correct information, and use simple checking and text-entry skills to improve accuracy.\nLearners will ask and answer short questions about contact details while practising confidence with Google tools and saved files.",
            "staging": "Session 1: 9:00-10:30 - onboard any new learners, check computer access, batteries, Google login, Google Drive, and Gmail. Early finishers use typing.com or Be Connected.\nSession 2: 10:45-12:00 - revisit digital contact list skills and demonstrate editing, correcting, and saving information in a spreadsheet or document.\nSession 3: 12:30-13:45 - learners practise entering, editing, and checking contact records with trainer support, followed by short oral review and Kahoot recap if time permits.\nDifferentiation / support: repeated demonstrations, paired support, simplified editing targets, and extension through extra records or categories.\nAssessment / evidence: edited contact list entries, observed error correction, successful reopening and saving of files, and teacher notes on learner confidence.",
            "resources": "Google account, Google Drive, and Gmail\nAddress book/contact list examples\nSpreadsheet or table template\ntyping.com\nBe Connected website\nKahoot\nComputers or laptops",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Friday",
            "date": "30 / 01 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Review and consolidation: contacts, Google access, and spreadsheet practice",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment",
            "objectives": "By the end of the day, learners will be able to review the week’s contact list and spreadsheet skills, locate their saved files, explain simple spreadsheet features, and demonstrate improved confidence with Google login, Drive, and Gmail.\nLearners will consolidate digital literacy skills through guided review and low-pressure checking activities.",
            "staging": "Session 1: 9:00-10:30 - setup routine with computer access, Google login, Drive, Gmail, batteries, and independent typing.com or Be Connected practice for early finishers.\nSession 2: 10:45-12:00 - learners continue address book and spreadsheet practice, locate their saved work, review entries, and correct errors with teacher support.\nSession 3: 12:30-13:45 - class review of learning from the week, short explanation of spreadsheet features by learners, and Kahoot review to reinforce confidence and recall.\nDifferentiation / support: one-to-one troubleshooting, repeated login/navigation support, paired help, and simplified completion goals for learners still building confidence.\nAssessment / evidence: located saved file, reviewed and corrected spreadsheet work, oral explanation of simple features, and teacher notes on learner progress and support provided.",
            "resources": "Address-Book-Training_2026_01_28.xlsx\nGoogle Sheets or Excel\nGoogle Drive\ntyping.com\nBe Connected website\nKahoot\nComputers or laptops",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
    ]

    for col, day in enumerate(days):
        set_cell_text(t1.cell(0, col), f"{day['name']}\n{day['date']}")
        set_cell_text(t1.cell(1, col), day["mode"])
        set_cell_text(t1.cell(2, col), "Synchronous")
        set_cell_text(t1.cell(3, col), "Topic:")
        set_cell_text(t1.cell(4, col), day["topic"])
        set_cell_text(t1.cell(5, col), "Units of Competency:")
        set_cell_text(t1.cell(6, col), day["units"])
        set_cell_text(t1.cell(7, col), "Lesson objectives:")
        set_cell_text(t1.cell(8, col), day["objectives"])
        set_cell_text(t1.cell(9, col), "Staging/Activities:")
        set_cell_text(t1.cell(10, col), day["staging"])

        set_cell_text(t2.cell(0, col), "Resources:")
        set_cell_text(t2.cell(1, col), day["resources"])
        set_cell_text(t2.cell(2, col), "Asynchronous")
        set_cell_text(t2.cell(3, col), "Topic:")
        set_cell_text(t2.cell(4, col), day["async_topic"])
        set_cell_text(t2.cell(5, col), "Units of Competency:")
        set_cell_text(t2.cell(6, col), day["async_units"])
        set_cell_text(t2.cell(7, col), "Lesson objectives:")
        set_cell_text(t2.cell(8, col), day["async_objectives"])
        set_cell_text(t2.cell(9, col), "Staging/Activities:")
        set_cell_text(t2.cell(10, col), day["async_staging"])
        set_cell_text(t2.cell(11, col), "Resources:")
        set_cell_text(t2.cell(12, col), day["async_resources"])

    set_cell_text(
        t3.cell(1, 0),
        "Provide additional support with login, browser navigation, Google Drive, Gmail, spreadsheet navigation, typing, and saving files as required.\n\n"
        "Pair confident and less confident learners for spreadsheet and contact-list practice where useful.\n\n"
        "Use short instructions, repeated demonstrations, and accuracy checks for learners who need extra support with entering and correcting digital information.\n\n"
        "Offer extension tasks for faster learners through extra contact entries, improved categorisation, or more independent spreadsheet work.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
