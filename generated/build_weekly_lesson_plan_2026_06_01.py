from copy import deepcopy
from pathlib import Path

from docx import Document


BASE = Path("/Users/simonjudge/Documents/MTC")
TEMPLATE = BASE / "01_Lesson_Plans/lesson_plans/templates/ET F419 - Weekly Lesson Plan Master.docx"
OUTPUT = BASE / "01_Lesson_Plans/lesson_plans/examples/ET_F419_Lesson_Plan_2026-06-01_to_2026-06-05_RED-01_Pathway.docx"


def set_cell_text(cell, text):
    cell.text = text


def main():
    doc = Document(str(TEMPLATE))

    t0, t1, t2, t3 = doc.tables

    set_cell_text(t0.cell(2, 1), "Redfern")
    set_cell_text(t0.cell(2, 3), "RED-01 Pathway")
    set_cell_text(t0.cell(3, 1), "01/06/2026")
    set_cell_text(t0.cell(3, 3), "05/06/2026")
    set_cell_text(t0.cell(4, 1), "Simon Judge")
    set_cell_text(t0.cell(5, 1), "Only as needed during delivery.")
    set_cell_text(
        t0.cell(5, 3),
        "Use repeated modelling, file-management support, calm redirection, smartphone access support, and extension tasks as required.",
    )
    set_cell_text(
        t0.cell(6, 1),
        "☐ Mainstream SEE\n☒ Mainstream SEE (Canvas)\n☒ Workforce Preparation\n☐ Co-delivery (e.g. Business)\n☐ Industry Pathway (e.g. Pathway to Individual Care)\n☐ Community Access (e.g. Citizenship)\n☐ Distance\n☐ Other:",
    )
    set_cell_text(
        t0.cell(7, 1),
        "☐ 22637VIC Course in English as an Additional Language (Course in EAL)\n☒ 22688VIC Course in Initial General Education for Adults (CGEA Initial)\n☒ 22689VIC Certificate I General Education for Adults (Introductory) (CGEA I Intro)\n☐ FSK20119 Certificate II in Skills for Work and Vocational Pathways (FSK II)\n☐ Other:",
    )

    days = [
        {
            "name": "Monday",
            "date": "01 / 06 / 2026",
            "mode": "Synchronous  ☐\nAsynchronous  ☐\nNo Class Scheduled  ☒",
            "topic": "No class scheduled / planning and preparation day",
            "units": "N/A - no scheduled class.",
            "objectives": "No scheduled class. Preparation only.",
            "staging": "Prepare differentiated activities for Tuesday to Friday.\nReview student notes and support needs.\nOrganise digital literacy, Canvas, file-management, and workplace communication resources.",
            "resources": "Weekly preparation notes\nExisting student notes\nPrepared digital literacy activity files",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Tuesday",
            "date": "02 / 06 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Digital creativity, safe AI use, and introductory poster creation",
            "units": "VU22353 Use basic oral communication skills\nVU22349 Create short simple texts for learning purposes\nVU22350 Engage with short simple texts for learning purposes\nDigital literacy / workforce preparation support",
            "objectives": "By the end of the day, learners will be able to describe one safe use of AI, contribute simple ideas for a poster, and participate in a supported digital creation task.\nLearners needing more support will complete a simple title-and-image task.\nStronger learners will add extra text or independent research.",
            "staging": "Session 1: 9:00-10:30 - warm-up discussion about digital tools learners already know; teacher models safe AI use and simple poster planning.\nSession 2: 10:45-12:00 - guided poster creation in pairs or small group with shared-screen modelling and oral language support.\nSession 3: 12:30-13:45 - saving, naming, and reviewing work; short whole-class reflection or low-pressure knowledge check.\nDifferentiation / support: sentence starters, visual modelling, paired work, repeated save/name/find routine, extension for fast finishers.\nAssessment / evidence: teacher observation of participation, saved draft files, short oral responses, and notes on independent device use.",
            "resources": "Computers or laptops\nInternet access\nGoogle Gemini poster demonstration\nMClient Wi-Fi reference if needed\nTeacher modelling on shared screen",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Wednesday",
            "date": "03 / 06 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Writing a clear workplace email and completing a simple digital workplace task",
            "units": "VU22349 Create short simple texts for learning purposes\nVU22350 Engage with short simple texts for learning purposes\nBSBTEC201 Use business software applications\nFSKDIG002 digital workplace task support",
            "objectives": "By the end of the day, learners will be able to identify the parts of a workplace email, write a short email for a practical class scenario, and complete a simple workplace-style digital task.\nLearners will choose the correct app for a task and save work using a clear filename.",
            "staging": "Session 1: 9:00-10:30 - introduce workplace email structure, key vocabulary, and model examples; complete matching and sequencing tasks.\nSession 2: 10:45-12:00 - guided writing of a short workplace email and teacher demonstration of choosing the right app and saving files.\nSession 3: 12:30-13:45 - mixed-level workplace task completion using either the business software sheet or the FSKDIG002 task, followed by review of saved files and email language.\nDifferentiation / support: word banks, model email on screen, paired support, repeated file-management routine, extension through longer digital workplace task.\nAssessment / evidence: completed activity sheets, saved document or folder, short email draft, teacher observation of file naming and app choice.",
            "resources": "Digital_Literacy_Email_Activity_Sheet_with_Crossword.docx\nActivity Sheet - Choose the Right App, Save the File, and Send It.docx\nFSKDIG002_Advanced_Digital_Workplace_Task_Activity_Sheet.docx\nComputers or laptops\nEmail platform",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Thursday",
            "date": "04 / 06 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Reading Australian websites, checking usefulness and trustworthiness, and guided digital participation",
            "units": "VU22350 Engage with short simple texts for learning purposes\nVU22353 Use basic oral communication skills\nACSF digital literacy practice\nFormative website-reading and comparison support",
            "objectives": "By the end of the day, learners will be able to identify key website features, skim and scan for important information, and make a simple comment about website purpose or usefulness.\nLearners will complete a website task at beginner, ACSF 2, or ACSF 3 level with support matched to need.",
            "staging": "Session 1: 9:00-10:30 - teacher-led watch / do together / say it together demonstration using a shared website and explicit teaching of headings, links, buttons, menus, and URLs.\nSession 2: 10:45-12:00 - differentiated website tasks: beginner website information sheet, ACSF 2 single-site activity, or ACSF 3 comparison activity.\nSession 3: 12:30-13:45 - supported completion, oral review, and short class check-in or Kahoot focused on website purpose, audience, and trustworthiness.\nDifferentiation / support: levelled activity sheets, shared reading, guided copying, oral rehearsal before writing, calm redirection, and extension through two-site comparison.\nAssessment / evidence: completed website sheets, filed PDFs, teacher notes on navigation and comprehension, oral responses about usefulness and trustworthiness.",
            "resources": "Beginner Digital Literacy Activity.docx\nACSF Level 2 Digital Literacy Activity Task.docx\nACSF Level 3 Digital Literacy Activity Task.docx\nFiled website activity PDFs from 2026-06-04\nComputers or laptops\nInternet access",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Friday",
            "date": "05 / 06 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Canvas access on smartphones, digital safety, and business software practice",
            "units": "BSBTEC201 Use business software applications\nVU22349 Create short simple texts for learning purposes\nVU22350 Engage with short simple texts for learning purposes\nDigital literacy and Canvas access support",
            "objectives": "By the end of the day, learners will be able to attempt Canvas access on a smartphone while following privacy rules, identify the correct app or tool for a digital task, and create, save, or send a simple workplace-style document.\nLearners will build confidence with troubleshooting and Canvas access support.",
            "staging": "Session 1: 9:00-10:30 - review privacy rules and model how to check Canvas access safely on a smartphone without exposing passwords.\nSession 2: 10:45-12:00 - learners complete the Canvas smartphone worksheet with trainer support; low-pressure troubleshooting is provided for login, small-screen use, or saved-password checks.\nSession 3: 12:30-13:45 - learners complete the business software / digital safety document task, including choosing the right app, saving files clearly, and sending or preparing work for review.\nDifferentiation / support: one-to-one device support, screen modelling, simplified checklist steps, extension through fuller document creation or more independent email completion.\nAssessment / evidence: completed smartphone access worksheet, completed business software activity, teacher notes on login success or barriers, saved files, and observation of privacy-safe behaviour.",
            "resources": "RED-01_Canvas_Phone_Access_Check_Student_Worksheet_2026-06-05.docx\nRED-01_Canvas_Smartphone_Access_Student_Worksheet_2026-06-05.docx\nActivity Sheet - Choose the Right App, Save the File, and Send It.docx\n2026-06-05_MTC_Redfern_RED-01_Pathway_ACSF2_Digital-Literacy-Activity_BSBTEC201.md\nComputers or laptops with internet access\nSmartphones for learners who choose to test phone access",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
    ]

    for col, day in enumerate(days):
        set_cell_text(t1.cell(0, col), f"{day['name']}\n{day['date']}")
        set_cell_text(t1.cell(1, col), day["mode"])
        set_cell_text(t1.cell(2, col), "Synchronous")
        set_cell_text(t1.cell(3, col), "Topic:")
        set_cell_text(t1.cell(4, col), day["topic"])
        set_cell_text(t1.cell(5, col), "Units of Competency:")
        set_cell_text(t1.cell(6, col), day["units"])
        set_cell_text(t1.cell(7, col), "Lesson objectives:")
        set_cell_text(t1.cell(8, col), day["objectives"])
        set_cell_text(t1.cell(9, col), "Staging/Activities:")
        set_cell_text(t1.cell(10, col), day["staging"])

        set_cell_text(t2.cell(0, col), "Resources:")
        set_cell_text(t2.cell(1, col), day["resources"])
        set_cell_text(t2.cell(2, col), "Asynchronous")
        set_cell_text(t2.cell(3, col), "Topic:")
        set_cell_text(t2.cell(4, col), day["async_topic"])
        set_cell_text(t2.cell(5, col), "Units of Competency:")
        set_cell_text(t2.cell(6, col), day["async_units"])
        set_cell_text(t2.cell(7, col), "Lesson objectives:")
        set_cell_text(t2.cell(8, col), day["async_objectives"])
        set_cell_text(t2.cell(9, col), "Staging/Activities:")
        set_cell_text(t2.cell(10, col), day["async_staging"])
        set_cell_text(t2.cell(11, col), "Resources:")
        set_cell_text(t2.cell(12, col), day["async_resources"])

    set_cell_text(
        t3.cell(1, 0),
        "Mei Lee Lim - repeated support with saving, naming, locating, and reopening files using one consistent routine.\n\n"
        "Dryl McGill - low-pressure troubleshooting, supported step-by-step digital tasks, and brief check-ins due to low frustration tolerance and low engagement.\n\n"
        "Sean Reilly - smartphone Canvas access support; larger screen or trainer-assisted navigation may be needed.\n\n"
        "Mark Williams, Terence Walker, and Djordje Milutinovich - protected time and trainer support for progressive assessments.\n\n"
        "Paul Martin - provide extension activities if he completes core tasks quickly.\n\n"
        "Use calm redirection and structured, concrete tasks for any learner who becomes unsettled or disengaged.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
