from pathlib import Path

from docx import Document


BASE = Path("/Users/simonjudge/Documents/MTC")
TEMPLATE = BASE / "01_Lesson_Plans/lesson_plans/templates/ET F419 - Weekly Lesson Plan Master.docx"
OUTPUT = BASE / "01_Lesson_Plans/lesson_plans/examples/ET_F419_Lesson_Plan_2026_02_02_to_2026_02_06.docx"


def set_cell_text(cell, text):
    cell.text = text


def main():
    doc = Document(str(TEMPLATE))
    t0, t1, t2, t3 = doc.tables

    set_cell_text(t0.cell(2, 1), "Redfern")
    set_cell_text(t0.cell(2, 3), "RED-01 Pathway")
    set_cell_text(t0.cell(3, 1), "02/02/2026")
    set_cell_text(t0.cell(3, 3), "06/02/2026")
    set_cell_text(t0.cell(4, 1), "Simon Judge")
    set_cell_text(t0.cell(5, 1), "Only as needed during delivery.")
    set_cell_text(
        t0.cell(5, 3),
        "Use staged modelling, repeated instructions, login support, reading support, spreadsheet support, and extension tasks where required.",
    )
    set_cell_text(
        t0.cell(6, 1),
        "☐ Mainstream SEE\n☒ Mainstream SEE (Canvas)\n☒ Workforce Preparation\n☐ Co-delivery (e.g. Business)\n☐ Industry Pathway (e.g. Pathway to Individual Care)\n☐ Community Access (e.g. Citizenship)\n☐ Distance\n☐ Other:",
    )
    set_cell_text(
        t0.cell(7, 1),
        "☐ 22637VIC Course in English as an Additional Language (Course in EAL)\n☒ 22688VIC Course in Initial General Education for Adults (CGEA Initial)\n☒ 22689VIC Certificate I General Education for Adults (Introductory) (CGEA I Intro)\n☐ FSK20119 Certificate II in Skills for Work and Vocational Pathways (FSK II)\n☐ Other:",
    )

    days = [
        {
            "name": "Monday",
            "date": "02 / 02 / 2026",
            "mode": "Synchronous  ☐\nAsynchronous  ☐\nNo Class Scheduled  ☒",
            "topic": "No class scheduled / preparation day",
            "units": "N/A - no scheduled class.",
            "objectives": "No scheduled class.",
            "staging": "Prepare digital literacy materials for Tuesday to Friday lessons.\nReview official letter examples, spreadsheet file, and differentiated support requirements.\nCheck device access, Google tools, and class resources.",
            "resources": "Prepared weekly resources\nOfficial letter examples\nExchange rate spreadsheet\nCanvas and Google tools",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Tuesday",
            "date": "03 / 02 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☒\nNo Class Scheduled  ☐",
            "topic": "Digital reading and document access: official letters and key information",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22350 Engage with short simple texts for learning purposes\nVU22349 Create short simple texts for learning purposes",
            "objectives": "By the end of the day, learners will be able to open a digital document, identify key information in an official letter, and record simple details such as name, date, organisation, and purpose.\nLearners will practise scrolling, locating headings, and reading short sections on screen with support.",
            "staging": "Session 1: 9:00-10:30 - teacher supports device login, opening documents, and locating key features in an official letter such as heading, date, recipient, and purpose.\nSession 2: 10:45-12:00 - learners read a mock letter on screen and complete guided tasks to identify important details and underline or record key information.\nSession 3: 12:30-13:45 - whole-class review of key information and short typed responses summarising what the letter is about.\nDifferentiation / support: read-aloud support, highlighted vocabulary, sentence starters, shared-screen modelling, and extension through extra detail or comparison of two letters.\nAssessment / evidence: successful opening of the document, identified key details, typed short responses, and teacher observation of on-screen reading skills.",
            "resources": "Mock Medical Letter Mary Smith 2026-02-05.docx\nComputers or laptops\nGoogle Drive or local file access\nProjector / shared screen\nTeacher comprehension prompts",
            "async_topic": "Jobs26 Canvas course - course overview and digital skills orientation",
            "async_units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22349 Create short simple texts for learning purposes",
            "async_objectives": "Learners will access the Jobs26 Canvas course independently, review the course syllabus, and complete selected orientation tasks such as Course Overview activities, discussion, and Digital Skills survey work.",
            "async_staging": "Students complete self-paced Canvas work in their own time. Focus areas may include Course Overview: Activity 1, Course Overview: Activity 2, Course Overview discussion, and Digital Skills: Activity 3 - Survey.\nTeacher checks completion and provides follow-up support in class where needed.",
            "async_resources": "Jobs26 Canvas course syllabus\nCourse Overview activities\nDigital Skills survey\nCanvas login details",
        },
        {
            "name": "Wednesday",
            "date": "04 / 02 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☒\nNo Class Scheduled  ☐",
            "topic": "Spreadsheet skills: exchange rates, data entry, and simple calculations",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment",
            "objectives": "By the end of the day, learners will be able to open a spreadsheet, identify labels and values, interpret a simple exchange-rate calculation, and enter or edit basic spreadsheet data with support.\nLearners will practise reading cells, checking values, and understanding a simple converted amount.",
            "staging": "Session 1: 9:00-10:30 - teacher models opening the exchange rates spreadsheet, identifying headings, and explaining amount, currency, exchange rate, and converted amount.\nSession 2: 10:45-12:00 - learners complete guided spreadsheet tasks by locating cells, changing simple values, and observing how the converted amount changes.\nSession 3: 12:30-13:45 - class review of spreadsheet terms, short oral questions, and independent or paired practice entering new amounts or currencies with teacher support.\nDifferentiation / support: one-step-at-a-time modelling, repeated explanation of rows, columns, and cells, paired help, and extension through extra examples or additional calculations.\nAssessment / evidence: opened spreadsheet, completed guided data-entry task, observed understanding of the converted amount, and teacher notes on digital confidence.",
            "resources": "Exchanges Rates_2026_02_04.xlsx\nExcel or Google Sheets\nComputers or laptops\nProjector / shared screen\nTeacher spreadsheet checklist",
            "async_topic": "Jobs26 Canvas course - cover letters and job advertisements",
            "async_units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22349 Create short simple texts for learning purposes\nVU22350 Engage with short simple texts for learning purposes",
            "async_objectives": "Learners will complete self-paced Canvas tasks related to reading job advertisements and responding to cover letter activities using supported digital literacy and text-entry skills.",
            "async_staging": "Students complete selected tasks in their own time, such as Cover Letters: Activities 1 to 3 and Job Advertisements: Activities 1 and 2.\nTeacher reviews progress and supports learners to catch up during live sessions where required.",
            "async_resources": "Jobs26 Canvas course syllabus\nCover Letters activities\nJob Advertisements activities\nCanvas login details",
        },
        {
            "name": "Thursday",
            "date": "05 / 02 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☒\nNo Class Scheduled  ☐",
            "topic": "Digital reading and comparison: action plans and My Aged Care letters",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22350 Engage with short simple texts for learning purposes\nVU22349 Create short simple texts for learning purposes",
            "objectives": "By the end of the day, learners will be able to open and compare two digital documents, identify action points and support contacts, and record simple information in a structured digital response.\nLearners will practise locating bullet points, contact numbers, and next steps in formal documents.",
            "staging": "Session 1: 9:00-10:30 - teacher introduces the Action Plan and My Aged Care letter, models how to compare purpose, headings, and important action points.\nSession 2: 10:45-12:00 - learners complete guided reading and comparison tasks, identifying key contacts, action steps, and support information in each document.\nSession 3: 12:30-13:45 - learners type short responses or fill in a comparison table summarising what each document is for and what action is needed.\nDifferentiation / support: highlighted sections, vocabulary support, paired reading, read-aloud modelling, and extension through fuller summaries or additional comparison points.\nAssessment / evidence: completed comparison task, identified action points and contacts, typed short summary, and teacher observation of document navigation and reading comprehension.",
            "resources": "Action_Plan_Vulnerable_2026_02_05.docx\nMock_My_Aged_Care_Assessment_Letter_Mary_Smith_2026_02_05.docx\nComputers or laptops\nComparison table or teacher worksheet\nProjector / shared screen",
            "async_topic": "Jobs26 Canvas course - job application forms and interviews",
            "async_units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22349 Create short simple texts for learning purposes\nVU22353 Use basic oral communication skills",
            "async_objectives": "Learners will work through self-paced Canvas tasks linked to job application forms and interview preparation, using digital navigation, reading, and short written response skills.",
            "async_staging": "Students complete selected Canvas tasks in their own time, including Job Application Forms: Activity 1 and Job Interviews: Activities 1 to 5 as appropriate.\nTeacher follows up in class with clarification, discussion, and technical support where needed.",
            "async_resources": "Jobs26 Canvas course syllabus\nJob Application Forms activity\nJob Interviews activities\nCanvas login details",
        },
        {
            "name": "Friday",
            "date": "06 / 02 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☒\nNo Class Scheduled  ☐",
            "topic": "Digital literacy consolidation: documents, spreadsheets, and guided review",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22349 Create short simple texts for learning purposes",
            "objectives": "By the end of the day, learners will be able to review the week’s digital tasks, reopen documents and spreadsheets, locate saved files, and complete a guided recap activity showing confidence with opening, reading, and editing digital materials.\nLearners will consolidate device login, file access, spreadsheet reading, and document-response skills.",
            "staging": "Session 1: 9:00-10:30 - review the week’s digital literacy tasks, including opening official letters, reading key information, and using the exchange-rate spreadsheet.\nSession 2: 10:45-12:00 - learners reopen saved files, revisit unfinished tasks, and complete a short recap activity using both document and spreadsheet skills.\nSession 3: 12:30-13:45 - teacher checks learner confidence with locating files, reading key details, and entering short responses; class ends with a simple quiz or Kahoot review if appropriate.\nDifferentiation / support: one-to-one troubleshooting, repeated login and file-location support, paired help, and simplified completion goals for learners needing more time.\nAssessment / evidence: reopened files successfully, completed recap activity, demonstrated basic file-location and response skills, and teacher notes on learner progress and support provided.",
            "resources": "Mock Medical Letter Mary Smith 2026-02-05.docx\nMock My Aged Care Assessment Letter Mary Smith 2026-02-05.docx\nExchanges Rates_2026_02_04.xlsx\nComputers or laptops\nGoogle Drive or local file access\nTeacher review checklist",
            "async_topic": "Jobs26 Canvas course - resumes, point of sale, and progress check",
            "async_units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22349 Create short simple texts for learning purposes",
            "async_objectives": "Learners will complete remaining self-paced Canvas tasks and review activities, including resume tasks, point of sale tasks, and the progress check survey, as part of approximately five hours of independent weekly course work.",
            "async_staging": "Students complete remaining Jobs26 tasks in their own time. This may include Resumes: Activities 1, 3, and 4; Point of Sale: Activities 1 to 3; Progress Check Survey; and other listed course tasks as appropriate.\nTeacher monitors progress and provides catch-up support during synchronous sessions.",
            "async_resources": "Jobs26 Canvas course syllabus\nResume activities\nPoint of Sale activities\nProgress Check Survey\nCanvas login details",
        },
    ]

    for col, day in enumerate(days):
        set_cell_text(t1.cell(0, col), f"{day['name']}\n{day['date']}")
        set_cell_text(t1.cell(1, col), day["mode"])
        set_cell_text(t1.cell(2, col), "Synchronous")
        set_cell_text(t1.cell(3, col), "Topic:")
        set_cell_text(t1.cell(4, col), day["topic"])
        set_cell_text(t1.cell(5, col), "Units of Competency:")
        set_cell_text(t1.cell(6, col), day["units"])
        set_cell_text(t1.cell(7, col), "Lesson objectives:")
        set_cell_text(t1.cell(8, col), day["objectives"])
        set_cell_text(t1.cell(9, col), "Staging/Activities:")
        set_cell_text(t1.cell(10, col), day["staging"])

        set_cell_text(t2.cell(0, col), "Resources:")
        set_cell_text(t2.cell(1, col), day["resources"])
        set_cell_text(t2.cell(2, col), "Asynchronous")
        set_cell_text(t2.cell(3, col), "Topic:")
        set_cell_text(t2.cell(4, col), day["async_topic"])
        set_cell_text(t2.cell(5, col), "Units of Competency:")
        set_cell_text(t2.cell(6, col), day["async_units"])
        set_cell_text(t2.cell(7, col), "Lesson objectives:")
        set_cell_text(t2.cell(8, col), day["async_objectives"])
        set_cell_text(t2.cell(9, col), "Staging/Activities:")
        set_cell_text(t2.cell(10, col), day["async_staging"])
        set_cell_text(t2.cell(11, col), "Resources:")
        set_cell_text(t2.cell(12, col), day["async_resources"])

    set_cell_text(
        t3.cell(1, 0),
        "Use short instructions, modelling, and repeated checks for learners needing support with device login, opening files, and reading on screen.\n\n"
        "Provide extra support with scrolling, locating headings, identifying key details, and entering simple written responses.\n\n"
        "Use paired reading, highlighted text, teacher read-aloud support, and simplified worksheets for learners needing LLN support.\n\n"
        "Offer extension tasks for faster learners through fuller summaries, extra spreadsheet examples, or more independent comparison work.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
