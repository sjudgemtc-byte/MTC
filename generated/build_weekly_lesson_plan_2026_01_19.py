from pathlib import Path

from docx import Document


BASE = Path("/Users/simonjudge/Documents/MTC")
TEMPLATE = BASE / "01_Lesson_Plans/lesson_plans/templates/ET F419 - Weekly Lesson Plan Master.docx"
OUTPUT = BASE / "01_Lesson_Plans/lesson_plans/examples/ET_F419_Lesson_Plan_2026_01_19_to_2026_01_23.docx"


def set_cell_text(cell, text):
    cell.text = text


def main():
    doc = Document(str(TEMPLATE))
    t0, t1, t2, t3 = doc.tables

    set_cell_text(t0.cell(2, 1), "Redfern")
    set_cell_text(t0.cell(2, 3), "RED-01 Pathway")
    set_cell_text(t0.cell(3, 1), "19/01/2026")
    set_cell_text(t0.cell(3, 3), "23/01/2026")
    set_cell_text(t0.cell(4, 1), "Simon Judge")
    set_cell_text(t0.cell(5, 1), "Only as needed during delivery.")
    set_cell_text(
        t0.cell(5, 3),
        "Use staged modelling, repeated instructions, file-management support, reading support, and extension tasks where required.",
    )
    set_cell_text(
        t0.cell(6, 1),
        "☐ Mainstream SEE\n☒ Mainstream SEE (Canvas)\n☒ Workforce Preparation\n☐ Co-delivery (e.g. Business)\n☐ Industry Pathway (e.g. Pathway to Individual Care)\n☐ Community Access (e.g. Citizenship)\n☐ Distance\n☐ Other:",
    )
    set_cell_text(
        t0.cell(7, 1),
        "☐ 22637VIC Course in English as an Additional Language (Course in EAL)\n☒ 22688VIC Course in Initial General Education for Adults (CGEA Initial)\n☒ 22689VIC Certificate I General Education for Adults (Introductory) (CGEA I Intro)\n☐ FSK20119 Certificate II in Skills for Work and Vocational Pathways (FSK II)\n☐ Other:",
    )

    days = [
        {
            "name": "Monday",
            "date": "19 / 01 / 2026",
            "mode": "Synchronous  ☐\nAsynchronous  ☐\nNo Class Scheduled  ☒",
            "topic": "No class scheduled / preparation day",
            "units": "N/A - no scheduled class.",
            "objectives": "No scheduled class. Preparation only.",
            "staging": "Prepare materials for Tuesday to Friday lessons.\nReview previous week's progression in past tense, digital literacy, and Canvas use.\nSet up differentiated support and extension activities.",
            "resources": "Previous week lesson plan\nPrepared classroom resources\nExisting digital literacy materials",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Tuesday",
            "date": "20 / 01 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Digital literacy foundation: logging in, opening class files, and typing simple routine sentences",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22349 Create short simple texts for learning purposes\nVU22353 Use basic oral communication skills",
            "objectives": "By the end of the day, learners will be able to log in to classroom devices and Google tools with support, open a class file, type simple sentences about everyday routines, and respond verbally to teacher prompts using routine vocabulary.\nLearners will practise basic keyboard use, on-screen navigation, and supported sentence entry.",
            "staging": "Session 1: 9:00-10:30 - teacher supports device login, opening browsers, and accessing the required class document; warm-up on routine vocabulary and oral question-and-answer practice.\nSession 2: 10:45-12:00 - learners follow step-by-step instructions to open a class file in Google Docs or a similar platform, read short prompts, and type simple answers about wake-up time, meals, travel, and bedtime.\nSession 3: 12:30-13:45 - learners review typed responses, practise simple editing, and save or re-open their work with teacher support.\nDifferentiation / support: visual prompts, oral rehearsal before typing, board modelling, one-to-one login support, and extension through added detail or extra typed sentences.\nAssessment / evidence: successful device and file access, typed routine sentences, observed keyboard and navigation skills, and teacher observation of supported oral responses.",
            "resources": "Computers or laptops\nGoogle Docs or classroom word-processing file\nTeacher-prepared daily routine prompts\nWhiteboard and markers\nInternet access",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Wednesday",
            "date": "21 / 01 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Digital reading and response task: Sleep and Daily Habits in Google Docs",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22349 Create short simple texts for learning purposes\nVU22350 Engage with short simple texts for learning purposes",
            "objectives": "By the end of the day, learners will be able to open and read the Sleep and Daily Habits document on screen, answer short comprehension questions by typing responses, and use simple digital reading strategies such as scrolling, locating questions, and following on-screen prompts.\nLearners will identify key vocabulary linked to sleep, health, and routines.",
            "staging": "Session 1: 9:00-10:30 - teacher models how to locate and open the Sleep and Daily Habits Google Doc, introduces key vocabulary, and demonstrates how to move through the document on screen.\nSession 2: 10:45-12:00 - learners complete guided digital reading of the document, using scrolling, teacher-led comprehension checks, and supported discussion.\nSession 3: 12:30-13:45 - learners type short answers or reflections in response to the text and review their work with teacher support.\nDifferentiation / support: read aloud support, vocabulary matching, sentence starters, paired navigation help, and extension through extra typed detail or comparison tasks.\nAssessment / evidence: successful opening of the Google Doc, typed comprehension responses, observed scrolling and navigation, and teacher observation of reading engagement.",
            "resources": "Sleep and Daily Habits Google Doc\nComputers or laptops\nGoogle Drive access\nVocabulary list\nProjected teacher model",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Thursday",
            "date": "22 / 01 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Google Docs basics: formatting, bold text, and following digital instructions",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22350 Engage with short simple texts for learning purposes",
            "objectives": "By the end of the day, learners will be able to open a Google Doc, follow simple written instructions, and use basic formatting tools such as bold text, headings, and simple text changes with support.\nLearners will build confidence navigating menus and toolbar icons in Google Docs.",
            "staging": "Session 1: 9:00-10:30 - teacher models how to open a Google Doc, identify the toolbar, and use bold text and other simple formatting features.\nSession 2: 10:45-12:00 - learners work through the Google Doc Intructions sheet step by step, completing formatting actions with shared-screen modelling and individual support.\nSession 3: 12:30-13:45 - learners repeat the skills more independently, save their work, and review file naming and locating documents in Drive.\nDifferentiation / support: repeated demonstration, one-step-at-a-time instructions, peer support, screenshot-style cues, and extension through extra formatting such as headings or alignment.\nAssessment / evidence: learner completion of formatting steps, saved document, observed navigation of menus and toolbar, and teacher notes on independence.",
            "resources": "Google Doc Intructions Google Doc\nGoogle Docs\nGoogle Drive access\nComputers or laptops\nProjector / shared screen",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Friday",
            "date": "23 / 01 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☐\nNo Class Scheduled  ☐",
            "topic": "Digital literacy consolidation: Canvas-supported review, Google Docs practice, and catch-up support",
            "units": "BSBTEC202 Use digital technologies to communicate in a work environment\nVU22349 Create short simple texts for learning purposes\nVU22353 Use basic oral communication skills",
            "objectives": "By the end of the day, learners will be able to review the week's digital tasks, access class content through Canvas with support, and complete or revisit key work in Google Docs and Google Drive.\nLearners requiring support will receive catch-up guidance to help them recover missed digital literacy tasks.",
            "staging": "Session 1: 9:00-10:30 - review the week's digital literacy content, including opening files, typing responses, reading on screen, and Google Docs formatting skills.\nSession 2: 10:45-12:00 - learners access class materials in Canvas and revisit unfinished tasks with teacher guidance, including document editing, typed responses, and locating saved files.\nSession 3: 12:30-13:45 - individual catch-up support, file checking, Google Drive review, and attendance follow-up actions where needed for absent learners; class finishes with a brief reflection on digital skills practised this week.\nDifferentiation / support: one-to-one catch-up support, repeated login and navigation assistance, short instructions, paired help, and simplified completion goals for learners with interrupted attendance.\nAssessment / evidence: accessed Canvas content, completed or revised digital task, saved document in the correct location, and teacher notes on learner follow-up and support provided.",
            "resources": "RED-01 Canvas course\nSleep and Daily Habits Google Doc\nGoogle Doc Intructions Google Doc\nGoogle Drive\nComputers or laptops\nTeacher follow-up checklist",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
    ]

    for col, day in enumerate(days):
        set_cell_text(t1.cell(0, col), f"{day['name']}\n{day['date']}")
        set_cell_text(t1.cell(1, col), day["mode"])
        set_cell_text(t1.cell(2, col), "Synchronous")
        set_cell_text(t1.cell(3, col), "Topic:")
        set_cell_text(t1.cell(4, col), day["topic"])
        set_cell_text(t1.cell(5, col), "Units of Competency:")
        set_cell_text(t1.cell(6, col), day["units"])
        set_cell_text(t1.cell(7, col), "Lesson objectives:")
        set_cell_text(t1.cell(8, col), day["objectives"])
        set_cell_text(t1.cell(9, col), "Staging/Activities:")
        set_cell_text(t1.cell(10, col), day["staging"])

        set_cell_text(t2.cell(0, col), "Resources:")
        set_cell_text(t2.cell(1, col), day["resources"])
        set_cell_text(t2.cell(2, col), "Asynchronous")
        set_cell_text(t2.cell(3, col), "Topic:")
        set_cell_text(t2.cell(4, col), day["async_topic"])
        set_cell_text(t2.cell(5, col), "Units of Competency:")
        set_cell_text(t2.cell(6, col), day["async_units"])
        set_cell_text(t2.cell(7, col), "Lesson objectives:")
        set_cell_text(t2.cell(8, col), day["async_objectives"])
        set_cell_text(t2.cell(9, col), "Staging/Activities:")
        set_cell_text(t2.cell(10, col), day["async_staging"])
        set_cell_text(t2.cell(11, col), "Resources:")
        set_cell_text(t2.cell(12, col), day["async_resources"])

    set_cell_text(
        t3.cell(1, 0),
        "Use short instructions, modelling, and repeated checks for learners needing support with reading, writing, and digital navigation.\n\n"
        "Provide repeated support with Google login, Canvas access, opening documents, saving work, and locating files in Drive.\n\n"
        "Use sentence frames, vocabulary banks, read-aloud support, and pair rehearsal for learners needing LLN support.\n\n"
        "Monitor attendance and confidence closely, and provide catch-up support and follow-up for learners who miss class time.\n\n"
        "Offer extension tasks for faster learners through added written detail, more independent formatting work, or extra digital practice.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
