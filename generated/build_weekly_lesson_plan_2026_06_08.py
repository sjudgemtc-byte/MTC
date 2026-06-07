from pathlib import Path

from docx import Document


BASE = Path("/Users/simonjudge/Documents/MTC")
TEMPLATE = BASE / "01_Lesson_Plans/lesson_plans/templates/ET F419 - Weekly Lesson Plan Master.docx"
OUTPUT = BASE / "01_Lesson_Plans/lesson_plans/examples/ET_F419_Lesson_Plan_2026_06_08_to_2026_06_12.docx"


def set_cell_text(cell, text):
    cell.text = text


def main():
    doc = Document(str(TEMPLATE))
    t0, t1, t2, t3 = doc.tables

    set_cell_text(t0.cell(2, 1), "Redfern")
    set_cell_text(t0.cell(2, 3), "RED-01 Pathway")
    set_cell_text(t0.cell(3, 1), "08/06/2026")
    set_cell_text(t0.cell(3, 3), "12/06/2026")
    set_cell_text(t0.cell(4, 1), "Simon Judge")
    set_cell_text(t0.cell(5, 1), "Only as needed during delivery.")
    set_cell_text(
        t0.cell(5, 3),
        "Use staged modelling, low-pressure troubleshooting, literacy support, repeated checking routines, device support, and extension tasks as needed.",
    )
    set_cell_text(
        t0.cell(6, 1),
        "☐ Mainstream SEE\n☒ Mainstream SEE (Canvas)\n☒ Workforce Preparation\n☐ Co-delivery (e.g. Business)\n☐ Industry Pathway (e.g. Pathway to Individual Care)\n☐ Community Access (e.g. Citizenship)\n☐ Distance\n☐ Other:",
    )
    set_cell_text(
        t0.cell(7, 1),
        "☐ 22637VIC Course in English as an Additional Language (Course in EAL)\n☒ 22688VIC Course in Initial General Education for Adults (CGEA Initial)\n☒ 22689VIC Certificate I General Education for Adults (Introductory) (CGEA I Intro)\n☐ FSK20119 Certificate II in Skills for Work and Vocational Pathways (FSK II)\n☐ Other:",
    )

    days = [
        {
            "name": "Monday",
            "date": "08 / 06 / 2026",
            "mode": "Synchronous  ☐\nAsynchronous  ☐\nNo Class Scheduled  ☒",
            "topic": "No class scheduled / planning and preparation day",
            "units": "N/A - no scheduled class.",
            "objectives": "No scheduled class. Preparation only.",
            "staging": "Prepare class files, check Canvas modules, and organise support materials for Tuesday to Friday delivery.",
            "resources": "Canvas notes\nPrepared class files\nPlanning notes",
            "async_topic": "N/A",
            "async_units": "N/A",
            "async_objectives": "N/A",
            "async_staging": "N/A",
            "async_resources": "N/A",
        },
        {
            "name": "Tuesday",
            "date": "09 / 06 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☒\nNo Class Scheduled  ☐",
            "topic": "Transport and Logistics course orientation, Canvas navigation, and digital typing practice",
            "units": "BSBTEC201 Business software applications\nVU22350 Short texts for learning\nVU22353 Basic oral communication",
            "objectives": "Learners open the RED-01 Canvas course, locate key weekly modules, and record simple information about their typing practice.",
            "staging": "Session 1: open Canvas, locate the weekly page, and identify key modules.\nSession 2: complete supported TypingClub practice and record time and activity.\nSession 3: complete a short transport vocabulary warm-up and save the file with a clear name.\nSupport: teacher modelling, paired login help, simplified record sheet, and extension typing for faster finishers.\nEvidence: Canvas access, typing record, saved file, and teacher observation.",
            "resources": "Canvas notes\nComputers or laptops\nCanvas access\nTypingClub link\nProgress note file",
            "async_topic": "Canvas typing",
            "async_units": "BSBTEC201 Software applications\nVU22349 Short texts",
            "async_objectives": "Students complete the typing task and record practice completed.",
            "async_staging": "Students practise touch typing for about 15 minutes and note what they practised.",
            "async_resources": "Canvas Digital Skills page\nTypingClub link\nLearner progress note",
        },
        {
            "name": "Wednesday",
            "date": "10 / 06 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☒\nNo Class Scheduled  ☐",
            "topic": "Transport and Logistics data entry using workplace-style tables",
            "units": "BSBTEC201 Business software applications\nVU22349 Short simple texts\nVU22350 Short texts for learning",
            "objectives": "Learners copy workplace-style transport and logistics data into a table, then check spelling, numbers, times, and pay rates carefully.",
            "staging": "Session 1: introduce the task, model the first row, and help learners choose a suitable program.\nSession 2: complete the main copying task using the full sheet or simple template.\nSession 3: check the table, identify jobs needing a licence, and use the approved AI prompt only after the first version is complete.\nSupport: simple template, one-to-one prompting, shared modelling, paired checking, and extension through job comparisons.\nEvidence: copied table, checklist, reflection, saved file, and teacher observation.",
            "resources": "Data entry sheet\nSimple template\nTrainer script\nComputers or laptops\nApproved AI prompt",
            "async_topic": "Canvas Word Processing",
            "async_units": "BSBTEC201 Software applications\nVU22349 Short texts",
            "async_objectives": "Students complete the Word Processing task and continue module work.",
            "async_staging": "Students continue the Word Processing tasks and note what they learned.",
            "async_resources": "Canvas Digital Skills page\nWord Processing module\nLearner progress note",
        },
        {
            "name": "Thursday",
            "date": "11 / 06 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☒\nNo Class Scheduled  ☐",
            "topic": "Transport and Logistics course orientation, Canvas navigation, and digital typing practice",
            "units": "BSBTEC201 Business software applications\nVU22350 Short texts for learning\nVU22353 Basic oral communication",
            "objectives": "Learners open the RED-01 Canvas course, locate key weekly modules, and record simple information about their typing practice.",
            "staging": "Session 1: open Canvas, locate the weekly page, and identify key modules.\nSession 2: complete supported TypingClub practice and record time and activity.\nSession 3: complete a short transport vocabulary warm-up and save the file with a clear name.\nSupport: teacher modelling, paired login help, simplified record sheet, and extension typing for faster finishers.\nEvidence: Canvas access, typing record, saved file, and teacher observation.",
            "resources": "Canvas notes\nComputers or laptops\nCanvas access\nTypingClub link\nProgress note file",
            "async_topic": "Canvas typing",
            "async_units": "BSBTEC201 Software applications\nVU22349 Short texts",
            "async_objectives": "Students complete the typing task and record practice completed.",
            "async_staging": "Students practise touch typing for about 15 minutes and note what they practised.",
            "async_resources": "Canvas Digital Skills page\nTypingClub link\nLearner progress note",
        },
        {
            "name": "Friday",
            "date": "12 / 06 / 2026",
            "mode": "Synchronous  ☒\nAsynchronous  ☒\nNo Class Scheduled  ☐",
            "topic": "Transport and Logistics data entry using workplace-style tables",
            "units": "BSBTEC201 Business software applications\nVU22349 Short simple texts\nVU22350 Short texts for learning",
            "objectives": "Learners copy workplace-style transport and logistics data into a table, then check spelling, numbers, times, and pay rates carefully.",
            "staging": "Session 1: introduce the task, model the first row, and help learners choose a suitable program.\nSession 2: complete the main copying task using the full sheet or simple template.\nSession 3: check the table, identify jobs needing a licence, and use the approved AI prompt only after the first version is complete.\nSupport: simple template, one-to-one prompting, shared modelling, paired checking, and extension through job comparisons.\nEvidence: copied table, checklist, reflection, saved file, and teacher observation.",
            "resources": "Data entry sheet\nSimple template\nTrainer script\nComputers or laptops\nApproved AI prompt",
            "async_topic": "Canvas Word Processing",
            "async_units": "BSBTEC201 Software applications\nVU22349 Short texts",
            "async_objectives": "Students complete the Word Processing task and continue module work.",
            "async_staging": "Students continue the Word Processing tasks and note what they learned.",
            "async_resources": "Canvas Digital Skills page\nWord Processing module\nLearner progress note",
        },
    ]

    for col, day in enumerate(days):
        set_cell_text(t1.cell(0, col), f"{day['name']}\n{day['date']}")
        set_cell_text(t1.cell(1, col), day["mode"])
        set_cell_text(t1.cell(2, col), "Synchronous")
        set_cell_text(t1.cell(3, col), "Topic:")
        set_cell_text(t1.cell(4, col), day["topic"])
        set_cell_text(t1.cell(5, col), "Units of Competency:")
        set_cell_text(t1.cell(6, col), day["units"])
        set_cell_text(t1.cell(7, col), "Lesson objectives:")
        set_cell_text(t1.cell(8, col), day["objectives"])
        set_cell_text(t1.cell(9, col), "Staging/Activities:")
        set_cell_text(t1.cell(10, col), day["staging"])

        set_cell_text(t2.cell(0, col), "Resources:")
        set_cell_text(t2.cell(1, col), day["resources"])
        set_cell_text(t2.cell(2, col), "Asynchronous")
        set_cell_text(t2.cell(3, col), "Topic:")
        set_cell_text(t2.cell(4, col), day["async_topic"])
        set_cell_text(t2.cell(5, col), "Units of Competency:")
        set_cell_text(t2.cell(6, col), day["async_units"])
        set_cell_text(t2.cell(7, col), "Lesson objectives:")
        set_cell_text(t2.cell(8, col), day["async_objectives"])
        set_cell_text(t2.cell(9, col), "Staging/Activities:")
        set_cell_text(t2.cell(10, col), day["async_staging"])
        set_cell_text(t2.cell(11, col), "Resources:")
        set_cell_text(t2.cell(12, col), day["async_resources"])

    set_cell_text(
        t3.cell(1, 0),
        "Use step-by-step modelling, paired checking, and repeated save / find / reopen routines for learners needing extra digital support.\n\n"
        "Provide the simple transport and logistics template for learners who need fewer columns, less reading, or more scaffolded table work.\n\n"
        "Support learners with login, navigation, and progress-check routines in Canvas, especially where module items require viewing or submission.\n\n"
        "Use sentence starters, oral rehearsal, and guided vocabulary review for learners needing LLN support.\n\n"
        "Offer extension through extra typing practice, job comparisons, price comparisons, or independent search tasks for faster finishers.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
